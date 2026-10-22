"""Excel QC report."""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILLS = {
    "PASS": PatternFill("solid", fgColor="D9F2E3"),
    "WARN": PatternFill("solid", fgColor="FFF1C9"),
    "FAIL": PatternFill("solid", fgColor="F8D3D0"),
    "INFO": PatternFill("solid", fgColor="E3ECFA"),
}
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(wb, title, headers, rows, widths, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
        if status_col is not None:
            st = row[status_col].value
            if st in FILLS:
                row[status_col].fill = FILLS[st]
                row[status_col].font = Font(bold=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def build_report(res, qnr: list[dict], url: str, qnr_name: str, out_path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    order = {"FAIL": 0, "WARN": 1, "INFO": 2, "PASS": 3}
    findings = sorted(res.findings, key=lambda f: (f.get("page") or 0, order.get(f["status"], 9)))
    f_counts = Counter(f["status"] for f in res.findings)
    v_counts = Counter(v["status"] for v in res.validations)
    not_seen = [q for q in qnr if q["id"] not in res.seen and q["type"] != "info"]

    # Summary
    ws = wb.create_sheet("Summary")
    rows = [
        ("Survey link QC report", ""),
        ("Live link", url),
        ("Questionnaire", qnr_name),
        ("Run at", datetime.now().strftime("%d %b %Y %H:%M")),
        ("End state", res.end_state),
        ("Pages visited", len(res.path)),
        ("QNR questions reached", f"{len(res.seen)} of {len([q for q in qnr if q['type'] != 'info'])}"),
        ("", ""),
        ("Text & option checks", "FAIL {FAIL} | WARN {WARN} | INFO {INFO} | PASS {PASS}".format(
            **{k: f_counts.get(k, 0) for k in FILLS})),
        ("Validation tests", "FAIL {FAIL} | WARN {WARN} | PASS {PASS}".format(
            **{k: v_counts.get(k, 0) for k in FILLS})),
        ("", ""),
        ("How to read", "FAIL = mismatch or validation missing. WARN = needs a human look (minor wording "
                        "difference, order change, unmatched question, no error message shown). INFO = expected "
                        "difference (e.g. randomised order). Questions not reached are usually routing: rerun "
                        "with 'answer' values set in the QNR JSON to walk other paths."),
    ]
    if res.end_text:
        rows.append(("End page text", res.end_text))
    if res.errors:
        rows.append(("Run errors", "\n".join(res.errors)))
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = WRAP
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110

    _sheet(wb, "Text & Options",
           ["Page", "QNR Q", "Link Q", "Check", "QNR (expected)", "Link (actual)", "Match %", "Status",
            "Note", "Matched by"],
           [[f.get("page"), f["question"], f.get("live_name", ""), f["check"], f["qnr"], f["live"],
             f["similarity"], f["status"], f["note"], f.get("matched_by", "")] for f in findings],
           [6, 10, 12, 18, 50, 50, 9, 9, 50, 14], status_col=7)

    _sheet(wb, "Validations",
           ["Page", "Question", "Test", "Input given", "Expected", "Result", "Message shown", "Status"],
           [[v["page"], v["question"], v["test"], v["input"], v["expected"], v["result"], v["message"],
             v["status"]] for v in res.validations],
           [6, 10, 32, 20, 16, 34, 60, 9], status_col=7)

    _sheet(wb, "Coverage",
           ["QNR Q", "Type", "Question text", "Reached on page", "Status", "Note"],
           [[q["id"], q["type"], q["text"], res.seen.get(q["id"], ""),
             "PASS" if q["id"] in res.seen else ("INFO" if q["type"] == "info" else "WARN"),
             "" if q["id"] in res.seen else ("Not reached on this path; check routing or test another path"
                                            + (f". QNR note: {q['notes']}" if q.get("notes") else ""))]
            for q in qnr],
           [10, 9, 70, 10, 9, 60], status_col=4)

    _sheet(wb, "Path log", ["Page", "Questions on page", "Answers given by bot", "Screenshot"],
           [[p["page"], p["questions"], p["answers"], p["screenshot"]] for p in res.path],
           [6, 40, 70, 30])

    wb.save(out_path)
    return out_path
